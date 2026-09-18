from __future__ import annotations

from pathlib import Path
from typing import Iterable, Sequence

from PIL import Image, ImageDraw, ImageFont
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_ALIGN_VERTICAL, WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(r"C:\Users\ADMIN\OneDrive\Desktop\synix\projects\preconstruction-ai")
WORK = ROOT / "work" / "understanding-guide"
OUTPUT = ROOT / "outputs" / "understanding-guide"
WORK.mkdir(parents=True, exist_ok=True)
OUTPUT.mkdir(parents=True, exist_ok=True)

DOCX_PATH = OUTPUT / "Structor_AI_Beginners_Guide.docx"
PROTOTYPE_IMAGE = ROOT / "references" / "structor-implementation-desktop-1440.png"
FLOW_IMAGE = WORK / "end_to_end_flow.png"
SHEET_IMAGE = WORK / "trade_sheet_anatomy.png"

# compact_reference_guide preset tokens
NAVY = "0B2545"
BLUE = "2E74B5"
DARK_BLUE = "1F4D78"
PALE_BLUE = "E8EEF5"
LIGHT_GRAY = "F2F4F7"
CALLOUT = "F4F6F9"
MUTED = "5B6573"
ORANGE = "F26B2B"  # Named override: Structor accent, used only for key labels.
GREEN = "2E7D32"
RED = "9B1C1C"
WHITE = "FFFFFF"
BLACK = "111827"


def rgb(hex_color: str) -> RGBColor:
    return RGBColor.from_string(hex_color)


def set_run_font(run, name="Calibri", size=None, color=None, bold=None, italic=None):
    run.font.name = name
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), name)
    if size is not None:
        run.font.size = Pt(size)
    if color is not None:
        run.font.color.rgb = rgb(color)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic


def set_cell_shading(cell, fill: str):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=80, start=120, bottom=80, end=120):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for edge, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        tag = "w:" + edge
        node = tc_mar.find(qn(tag))
        if node is None:
            node = OxmlElement(tag)
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def prevent_row_split(row):
    tr_pr = row._tr.get_or_add_trPr()
    cant_split = tr_pr.find(qn("w:cantSplit"))
    if cant_split is None:
        cant_split = OxmlElement("w:cantSplit")
        tr_pr.append(cant_split)


def set_table_borders(table, color="D7DCE2", size="6"):
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.first_child_found_in("w:tblBorders")
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = "w:" + edge
        node = borders.find(qn(tag))
        if node is None:
            node = OxmlElement(tag)
            borders.append(node)
        node.set(qn("w:val"), "single")
        node.set(qn("w:sz"), size)
        node.set(qn("w:space"), "0")
        node.set(qn("w:color"), color)


def set_table_geometry(table, widths_dxa: Sequence[int], indent_dxa=120):
    assert sum(widths_dxa) == 9360, widths_dxa
    table.autofit = False
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    tbl = table._tbl
    tbl_pr = tbl.tblPr

    tbl_w = tbl_pr.first_child_found_in("w:tblW")
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), "9360")
    tbl_w.set(qn("w:type"), "dxa")

    tbl_ind = tbl_pr.first_child_found_in("w:tblInd")
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), str(indent_dxa))
    tbl_ind.set(qn("w:type"), "dxa")

    old_grid = tbl.tblGrid
    for child in list(old_grid):
        old_grid.remove(child)
    for width in widths_dxa:
        grid_col = OxmlElement("w:gridCol")
        grid_col.set(qn("w:w"), str(width))
        old_grid.append(grid_col)

    for row in table.rows:
        for idx, cell in enumerate(row.cells):
            width = widths_dxa[min(idx, len(widths_dxa) - 1)]
            tc_pr = cell._tc.get_or_add_tcPr()
            tc_w = tc_pr.first_child_found_in("w:tcW")
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                tc_pr.append(tc_w)
            tc_w.set(qn("w:w"), str(width))
            tc_w.set(qn("w:type"), "dxa")
            set_cell_margins(cell)


def configure_table(table, widths_dxa: Sequence[int], header=True, compact=False):
    set_table_geometry(table, widths_dxa)
    set_table_borders(table)
    if header:
        set_repeat_table_header(table.rows[0])
    for r_idx, row in enumerate(table.rows):
        prevent_row_split(row)
        for cell in row.cells:
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            for p in cell.paragraphs:
                p.paragraph_format.space_before = Pt(0)
                p.paragraph_format.space_after = Pt(2 if compact else 3)
                p.paragraph_format.line_spacing = 1.08
                for run in p.runs:
                    set_run_font(run, size=9.2 if compact else 9.6, color=BLACK)
        if header and r_idx == 0:
            for cell in row.cells:
                set_cell_shading(cell, PALE_BLUE)
                for p in cell.paragraphs:
                    p.paragraph_format.keep_with_next = True
                    for run in p.runs:
                        set_run_font(run, size=9.2, color=NAVY, bold=True)


def add_table(doc, headers: Sequence[str], rows: Sequence[Sequence[str]], widths: Sequence[int], compact=False):
    table = doc.add_table(rows=1, cols=len(headers))
    for idx, value in enumerate(headers):
        table.rows[0].cells[idx].text = value
    for row_values in rows:
        cells = table.add_row().cells
        for idx, value in enumerate(row_values):
            cells[idx].text = str(value)
    configure_table(table, widths, header=True, compact=compact)
    after = doc.add_paragraph()
    after.paragraph_format.space_after = Pt(2)
    return table


def add_callout(doc, label: str, text: str, fill=CALLOUT, accent=ORANGE):
    table = doc.add_table(rows=1, cols=1)
    cell = table.cell(0, 0)
    cell.text = ""
    set_cell_shading(cell, fill)
    set_table_geometry(table, [9360])
    set_table_borders(table, color=accent, size="10")
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(0)
    r = p.add_run(label + "  ")
    set_run_font(r, size=10.5, color=accent, bold=True)
    r = p.add_run(text)
    set_run_font(r, size=10.5, color=BLACK)
    doc.add_paragraph().paragraph_format.space_after = Pt(1)


def add_body(doc, text: str, bold_lead: str | None = None):
    p = doc.add_paragraph()
    if bold_lead and text.startswith(bold_lead):
        r = p.add_run(bold_lead)
        set_run_font(r, bold=True, color=NAVY)
        r = p.add_run(text[len(bold_lead):])
        set_run_font(r, color=BLACK)
    else:
        r = p.add_run(text)
        set_run_font(r, color=BLACK)
    return p


def add_bullet(doc, text: str, level=0):
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Inches(0.375 + 0.25 * level)
    p.paragraph_format.first_line_indent = Inches(-0.188)
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.line_spacing = 1.25
    p_pr = p._p.get_or_add_pPr()
    num_pr = OxmlElement("w:numPr")
    ilvl = OxmlElement("w:ilvl")
    ilvl.set(qn("w:val"), str(level))
    num_id = OxmlElement("w:numId")
    num_id.set(qn("w:val"), "50")
    num_pr.append(ilvl)
    num_pr.append(num_id)
    p_pr.append(num_pr)
    r = p.add_run(text)
    set_run_font(r, color=BLACK)
    return p


def add_number(doc, text: str, level=0):
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Inches(0.375 + 0.25 * level)
    p.paragraph_format.first_line_indent = Inches(-0.188)
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.line_spacing = 1.25
    p_pr = p._p.get_or_add_pPr()
    num_pr = OxmlElement("w:numPr")
    ilvl = OxmlElement("w:ilvl")
    ilvl.set(qn("w:val"), str(level))
    num_id = OxmlElement("w:numId")
    num_id.set(qn("w:val"), "51")
    num_pr.append(ilvl)
    num_pr.append(num_id)
    p_pr.append(num_pr)
    r = p.add_run(text)
    set_run_font(r, color=BLACK)
    return p


def add_manual_number(doc, text: str, number: int, level=0):
    """Add a numbered paragraph whose sequence is explicitly controlled."""
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Inches(0.375 + 0.25 * level)
    p.paragraph_format.first_line_indent = Inches(-0.188)
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.line_spacing = 1.25
    r = p.add_run(f"{number}.  ")
    set_run_font(r, color=BLACK)
    r = p.add_run(text)
    set_run_font(r, color=BLACK)
    return p


def install_numbering(doc: Document):
    numbering = doc.part.numbering_part.element
    existing_abs = [int(x.get(qn("w:abstractNumId"))) for x in numbering.findall(qn("w:abstractNum"))]
    existing_num = [int(x.get(qn("w:numId"))) for x in numbering.findall(qn("w:num"))]
    abs_bullet = max(existing_abs or [0]) + 1
    abs_decimal = abs_bullet + 1

    def abstract(abs_id: int, fmt: str, text: str):
        a = OxmlElement("w:abstractNum")
        a.set(qn("w:abstractNumId"), str(abs_id))
        multi = OxmlElement("w:multiLevelType")
        multi.set(qn("w:val"), "multilevel")
        a.append(multi)
        for level in range(3):
            lvl = OxmlElement("w:lvl")
            lvl.set(qn("w:ilvl"), str(level))
            start = OxmlElement("w:start")
            start.set(qn("w:val"), "1")
            num_fmt = OxmlElement("w:numFmt")
            num_fmt.set(qn("w:val"), fmt)
            lvl_text = OxmlElement("w:lvlText")
            lvl_text.set(qn("w:val"), text if fmt == "bullet" else f"%{level + 1}.")
            suff = OxmlElement("w:suff")
            suff.set(qn("w:val"), "tab")
            p_pr = OxmlElement("w:pPr")
            tabs = OxmlElement("w:tabs")
            tab = OxmlElement("w:tab")
            tab.set(qn("w:val"), "num")
            tab.set(qn("w:pos"), str(540 + level * 360))
            tabs.append(tab)
            ind = OxmlElement("w:ind")
            ind.set(qn("w:left"), str(540 + level * 360))
            ind.set(qn("w:hanging"), "270")
            spacing = OxmlElement("w:spacing")
            spacing.set(qn("w:after"), "80")
            spacing.set(qn("w:line"), "300")
            spacing.set(qn("w:lineRule"), "auto")
            p_pr.extend([tabs, ind, spacing])
            lvl.extend([start, num_fmt, lvl_text, suff, p_pr])
            if fmt == "bullet":
                r_pr = OxmlElement("w:rPr")
                r_fonts = OxmlElement("w:rFonts")
                r_fonts.set(qn("w:ascii"), "Symbol")
                r_fonts.set(qn("w:hAnsi"), "Symbol")
                r_pr.append(r_fonts)
                lvl.append(r_pr)
            a.append(lvl)
        numbering.append(a)

    abstract(abs_bullet, "bullet", "\uf0b7")
    abstract(abs_decimal, "decimal", "%1.")

    for num_id, abs_id in ((50, abs_bullet), (51, abs_decimal)):
        if num_id in existing_num:
            continue
        n = OxmlElement("w:num")
        n.set(qn("w:numId"), str(num_id))
        aid = OxmlElement("w:abstractNumId")
        aid.set(qn("w:val"), str(abs_id))
        n.append(aid)
        numbering.append(n)


def add_page_number(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    r = paragraph.add_run("Page ")
    set_run_font(r, size=9, color=MUTED)
    fld_begin = OxmlElement("w:fldChar")
    fld_begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    fld_sep = OxmlElement("w:fldChar")
    fld_sep.set(qn("w:fldCharType"), "separate")
    fld_text = OxmlElement("w:t")
    fld_text.text = "1"
    fld_end = OxmlElement("w:fldChar")
    fld_end.set(qn("w:fldCharType"), "end")
    run = OxmlElement("w:r")
    run.extend([fld_begin, instr, fld_sep, fld_text, fld_end])
    paragraph._p.append(run)


def set_picture_alt(paragraph, description: str):
    drawings = paragraph._p.xpath(".//wp:docPr")
    for node in drawings:
        node.set("descr", description)


def create_diagrams():
    font_paths = [
        Path(r"C:\Windows\Fonts\calibri.ttf"),
        Path(r"C:\Windows\Fonts\arial.ttf"),
    ]
    font_path = next((p for p in font_paths if p.exists()), None)

    def font(size, bold=False):
        candidates = [Path(r"C:\Windows\Fonts\calibrib.ttf"), Path(r"C:\Windows\Fonts\arialbd.ttf")] if bold else font_paths
        fpath = next((p for p in candidates if p.exists()), font_path)
        return ImageFont.truetype(str(fpath), size) if fpath else ImageFont.load_default()

    # End-to-end flow
    img = Image.new("RGB", (1800, 760), "white")
    d = ImageDraw.Draw(img)
    title_font = font(42, True)
    box_font = font(27, True)
    small_font = font(22, False)
    d.text((65, 36), "The preconstruction data flow", fill="#0B2545", font=title_font)
    boxes = [
        ("1", "Design documents", "Drawings + specifications"),
        ("2", "Trade scopes", "What each subcontractor must include"),
        ("3", "Bid proposals", "Prices, exclusions, qualifications"),
        ("4", "Bid leveling", "Make every bid comparable"),
        ("5", "Budget + award", "Selected cost and contract scope"),
    ]
    x_positions = [55, 405, 755, 1105, 1455]
    for i, (num, heading, sub) in enumerate(boxes):
        x = x_positions[i]
        y = 205
        d.rounded_rectangle((x, y, x + 290, y + 235), radius=24, fill="#F4F6F9", outline="#D7DCE2", width=3)
        d.ellipse((x + 18, y + 18, x + 72, y + 72), fill="#F26B2B")
        d.text((x + 37, y + 27), num, fill="white", anchor="mm", font=font(26, True))
        d.multiline_text((x + 145, y + 104), heading, fill="#0B2545", anchor="mm", align="center", spacing=4, font=box_font)
        d.multiline_text((x + 145, y + 175), sub, fill="#5B6573", anchor="mm", align="center", spacing=4, font=small_font)
        if i < len(boxes) - 1:
            d.line((x + 295, y + 118, x_positions[i + 1] - 15, y + 118), fill="#2E74B5", width=6)
            d.polygon([(x_positions[i + 1] - 15, y + 118), (x_positions[i + 1] - 37, y + 106), (x_positions[i + 1] - 37, y + 130)], fill="#2E74B5")
    d.rounded_rectangle((360, 530, 1440, 675), radius=20, fill="#E8EEF5")
    d.text((900, 568), "Human approval is required at the important judgment points", fill="#1F4D78", anchor="mm", font=font(27, True))
    d.text((900, 620), "Approve scope evidence • confirm Y/N/P • approve adjustments • select bidder • approve contract", fill="#111827", anchor="mm", font=small_font)
    img.save(FLOW_IMAGE, quality=95)

    # Trade sheet anatomy
    img = Image.new("RGB", (1800, 1050), "white")
    d = ImageDraw.Draw(img)
    d.text((60, 40), "How one trade-level bid sheet is organized", fill="#0B2545", font=title_font)
    d.rounded_rectangle((70, 125, 1730, 245), radius=20, fill="#E8EEF5", outline="#B7C7D8", width=3)
    d.text((110, 155), "PROJECT / TRADE / ESTIMATOR", fill="#1F4D78", font=box_font)
    d.text((110, 202), "Example: Jelleff Recreation Center • 02 Demolition • assigned estimator", fill="#111827", font=small_font)

    d.rounded_rectangle((70, 285, 735, 815), radius=20, fill="#F4F6F9", outline="#D7DCE2", width=3)
    d.text((110, 320), "THE MASTER REQUIREMENT", fill="#F26B2B", font=box_font)
    left_lines = [
        "Each row is one scope item.",
        "Desired scope = what the GC wants included.",
        "Budget fields = quantity, unit and cost.",
        "Source reference = drawing/spec evidence.",
    ]
    y = 390
    for line in left_lines:
        d.ellipse((112, y + 7, 126, y + 21), fill="#2E74B5")
        d.text((145, y), line, fill="#111827", font=small_font)
        y += 86

    d.rounded_rectangle((780, 285, 1730, 815), radius=20, fill="#FFF7F2", outline="#F4B08B", width=3)
    d.text((820, 320), "REPEATING BIDDER GROUPS", fill="#F26B2B", font=box_font)
    bidder_boxes = [
        ("Bidder A", "Y/N/P", "Note", "Adjustment"),
        ("Bidder B", "Y/N/P", "Note", "Adjustment"),
        ("Bidder C", "Y/N/P", "Note", "Adjustment"),
    ]
    bx = 820
    for bidder, status, note, adj in bidder_boxes:
        d.rounded_rectangle((bx, 400, bx + 260, 700), radius=16, fill="white", outline="#D7DCE2", width=3)
        d.text((bx + 130, 438), bidder, fill="#0B2545", anchor="mm", font=font(25, True))
        d.text((bx + 130, 510), status, fill="#2E74B5", anchor="mm", font=small_font)
        d.text((bx + 130, 580), note, fill="#5B6573", anchor="mm", font=small_font)
        d.text((bx + 130, 650), adj, fill="#9B1C1C", anchor="mm", font=small_font)
        bx += 285

    d.rounded_rectangle((250, 865, 1550, 985), radius=20, fill="#0B2545")
    d.text((900, 905), "FORMULA OUTPUTS", fill="white", anchor="mm", font=box_font)
    d.text((900, 950), "Base bid + approved adjustments + bond = leveled total → rank", fill="white", anchor="mm", font=small_font)
    img.save(SHEET_IMAGE, quality=95)


def style_document(doc: Document):
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(1.0)
    section.bottom_margin = Inches(1.0)
    section.left_margin = Inches(1.0)
    section.right_margin = Inches(1.0)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)
    section.different_first_page_header_footer = True

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Calibri"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    normal.font.size = Pt(11)
    normal.font.color.rgb = rgb(BLACK)
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.25

    for style_name, size, color, before, after in (
        ("Heading 1", 16, BLUE, 18, 10),
        ("Heading 2", 13, BLUE, 14, 7),
        ("Heading 3", 12, DARK_BLUE, 10, 5),
    ):
        st = styles[style_name]
        st.font.name = "Calibri"
        st._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
        st._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
        st.font.size = Pt(size)
        st.font.bold = True
        st.font.color.rgb = rgb(color)
        st.paragraph_format.space_before = Pt(before)
        st.paragraph_format.space_after = Pt(after)
        st.paragraph_format.keep_with_next = True

    if "Figure Caption" not in styles:
        cap = styles.add_style("Figure Caption", WD_STYLE_TYPE.PARAGRAPH)
    else:
        cap = styles["Figure Caption"]
    cap.font.name = "Calibri"
    cap._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    cap._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    cap.font.size = Pt(9)
    cap.font.italic = True
    cap.font.color.rgb = rgb(MUTED)
    cap.paragraph_format.space_before = Pt(4)
    cap.paragraph_format.space_after = Pt(10)
    cap.paragraph_format.keep_with_next = False

    install_numbering(doc)

    # Running header/footer. No decorative title underline on the cover.
    header = section.header
    hp = header.paragraphs[0]
    hp.text = "STRUCTOR AI  |  BEGINNER'S GUIDE"
    hp.alignment = WD_ALIGN_PARAGRAPH.LEFT
    for run in hp.runs:
        set_run_font(run, size=8.5, color=MUTED, bold=True)
    footer = section.footer
    add_page_number(footer.paragraphs[0])


def add_heading(doc, text: str, level=1):
    p = doc.add_paragraph(text, style=f"Heading {level}")
    return p


def add_figure(doc, path: Path, caption: str, alt: str, width=6.45):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.keep_with_next = True
    p.add_run().add_picture(str(path), width=Inches(width))
    set_picture_alt(p, alt)
    cp = doc.add_paragraph(caption, style="Figure Caption")
    cp.alignment = WD_ALIGN_PARAGRAPH.CENTER


def page_break(doc):
    doc.add_page_break()


def build_document():
    create_diagrams()
    doc = Document()
    style_document(doc)

    # Editorial cover pattern
    for _ in range(4):
        doc.add_paragraph().paragraph_format.space_after = Pt(14)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("FIELD GUIDE")
    set_run_font(r, size=11, color=ORANGE, bold=True)
    p.paragraph_format.space_after = Pt(16)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(10)
    r = p.add_run("Understanding Structor AI")
    set_run_font(r, size=30, color=NAVY, bold=True)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(6)
    r = p.add_run("Marcel's preconstruction workflow, the Excel sheets, and the automation you promised")
    set_run_font(r, size=15, color=DARK_BLUE)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(40)
    r = p.add_run("A plain-language guide for Danial Khan / Synix Solutions")
    set_run_font(r, size=10.5, color=MUTED, italic=True)

    add_callout(
        doc,
        "THE ONE-SENTENCE VERSION",
        "Structor AI is meant to turn construction documents and subcontractor proposals into complete trade scopes, comparable bids, a defensible project budget, and eventually contract scope language.",
        fill=LIGHT_GRAY,
    )
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(36)
    r = p.add_run("Based on the 26 August 2026 discovery meeting and the Jelleff Recreation Center bid workbook")
    set_run_font(r, size=9.5, color=MUTED)

    page_break(doc)

    add_heading(doc, "How to use this guide", 1)
    add_body(doc, "Read Sections 1–4 first. They explain the business and workbook without assuming construction knowledge. Sections 5–8 explain the automation commitment and how the software should work. The final sections give you a practical checklist, questions for Marcel, and a glossary.")
    add_callout(doc, "MOST IMPORTANT DISTINCTION", "The Excel workbook is the current container for Marcel's decisions. The real product is the process that creates those decisions and moves them from drawings to scopes, bids, budgets, and contracts.")

    add_heading(doc, "Contents", 2)
    contents = [
        "1. The business in plain English",
        "2. What Marcel actually does every day",
        "3. The documents and construction vocabulary",
        "4. What every workbook sheet means",
        "5. A trade sheet explained cell by cell",
        "6. Bid leveling with a worked example",
        "7. What is manual, what Excel already automates, and what Structor should automate",
        "8. What you actually promised Marcel",
        "9. The recommended product milestones and data model",
        "10. How the AI components should work safely",
        "11. What you need from Marcel next",
        "12. Meeting questions and glossary",
    ]
    for item in contents:
        add_bullet(doc, item)

    page_break(doc)

    add_heading(doc, "1. The business in plain English", 1)
    add_body(doc, "Marcel works in preconstruction for a general contractor. Preconstruction is the planning, pricing, and purchasing work performed before physical construction begins. His company must understand a proposed building well enough to tell the owner what it will cost and to buy the work from specialist subcontractors without leaving gaps.")
    add_body(doc, "Think of the general contractor as the company assembling the entire building. It usually does not personally manufacture every door, pour every slab, install every pipe, and wire every room. It hires separate specialist companies, called subcontractors, for those pieces of work.")

    add_heading(doc, "The people involved", 2)
    add_table(
        doc,
        ["Party", "What that party does"],
        [
            ["Owner / client", "Pays for the project and defines the desired outcome."],
            ["Architect and engineers", "Create drawings and specifications describing the building."],
            ["General contractor (GC)", "Coordinates, prices, purchases, schedules, and manages the overall construction."],
            ["Bid captain", "Leads the preconstruction bid, sets deadlines, assigns trades, and assembles the final price."],
            ["Estimator / division leader", "Owns one or more trades, writes scopes, contacts bidders, levels bids, and recommends an award."],
            ["Subcontractor / sub", "A separate specialist company bidding to perform one trade, such as concrete or electrical work."],
        ],
        [2400, 6960],
    )

    add_heading(doc, "What Marcel is really selling", 2)
    add_body(doc, "He is not merely producing spreadsheets. He is producing confidence: confidence that the project scope is complete, that competing bids are comparable, that the budget does not hide missing work, and that the awarded contract clearly says what the subcontractor must provide.")
    add_figure(doc, FLOW_IMAGE, "Figure 1. The entire preconstruction flow in one picture.", "Five-stage flow from design documents through scopes, proposals, bid leveling, budget, and award, with human approval across all stages.")

    add_heading(doc, "The central risk", 2)
    add_body(doc, "A cheap-looking bid can be expensive if it excludes work that another bidder included. If nobody owns an item, the GC may discover the gap after the budget is committed. The missing work then becomes a cost overrun, dispute, or change order. Most of Marcel's process exists to prevent that outcome.")

    page_break(doc)

    add_heading(doc, "2. What Marcel actually does every day", 1)
    add_body(doc, "The meeting described the following operational sequence. In practice, several steps overlap and different estimators work in parallel.")
    steps = [
        ("Receive project documents", "The owner or architect provides drawings, specifications, reports, addenda, schedules, and instructions."),
        ("Perform a drawing flip", "The bid captain reviews the set, identifies which CSI divisions and trades apply, and assigns them to estimators."),
        ("Write one scope per trade", "Each estimator reads the relevant drawings and specifications and converts them into a checklist of everything that trade must include."),
        ("Invite subcontractors", "The estimator sends the package to several qualified companies through BuildingConnected, email, or another outreach process."),
        ("Answer questions and issue RFIs", "Subcontractors and estimators flag missing, unclear, or contradictory information before pricing is finalized."),
        ("Receive proposals", "Each bidder returns a price plus inclusions, exclusions, allowances, alternates, and qualifications, usually by email."),
        ("Level the bids", "The estimator compares each proposal against the same master scope, records Y/N/P, and adds internal adjustments for missing work."),
        ("Select a value and update the budget", "The reviewed bid becomes the current expected cost for that trade; all trade values roll into the project total."),
        ("Review scope and award", "The team meets with the preferred bidder, resolves remaining gaps, and makes an award recommendation."),
        ("Create Exhibit A", "The approved scope is rewritten or formatted as the variable scope section of the subcontract."),
    ]
    for title, detail in steps:
        p = add_number(doc, "")
        p.clear()
        r = p.add_run(title + ". ")
        set_run_font(r, color=NAVY, bold=True)
        r = p.add_run(detail)
        set_run_font(r, color=BLACK)

    add_heading(doc, "Why it takes so long", 2)
    add_body(doc, "Marcel said a drawing and specification package can contain roughly 1,000 to 5,000 pages. The team repeats the scope and leveling exercise across approximately 70 to 80 trades. The difficult work is not the addition formula; it is reading, classifying, comparing, and documenting thousands of facts and judgments.")

    add_heading(doc, "What Marcel personally contributes", 2)
    add_table(
        doc,
        ["Marcel's contribution", "Why software cannot blindly replace it"],
        [
            ["Trade ownership judgment", "A note can appear on one drawing but belong contractually to a different trade."],
            ["Scope completeness", "Generic scope language must be combined with project-specific details and coordination requirements."],
            ["Proposal interpretation", "Subcontractor wording is inconsistent and may imply partial coverage rather than a clean yes or no."],
            ["Adjustment pricing", "The cost of missing work depends on quantities, unit rates, market conditions, logistics, and risk."],
            ["Award recommendation", "Price matters, but so do capability, schedule, qualifications, relationships, and commercial risk."],
            ["Contract judgment", "An approved scope becomes a binding obligation and requires legal and operational review."],
        ],
        [3100, 6260],
    )

    page_break(doc)

    add_heading(doc, "3. The documents and construction vocabulary", 1)
    add_heading(doc, "Drawings versus specifications", 2)
    add_body(doc, "Drawings show where things are, their arrangement, dimensions, and relationships. The project manual or specifications explain materials, products, quality standards, execution requirements, testing, warranties, and other details. Marcel summarized the distinction as: drawings show what must be done; specifications explain how it must be done.")
    add_table(
        doc,
        ["Source", "Typical information", "Example"],
        [
            ["Architectural drawings", "Rooms, walls, doors, finishes, elevations, details", "Where gypsum partitions and flooring occur"],
            ["Structural drawings", "Foundations, slabs, beams, columns, steel and reinforcing", "Thickness and reinforcement of a concrete slab"],
            ["Civil drawings", "Grading, roads, exterior utilities, drainage and site work", "Sanitary line from the street to the building"],
            ["MEP/FP drawings", "Mechanical, electrical, plumbing and fire-protection systems", "Ductwork, power, piping and sprinkler layouts"],
            ["Specifications / project manual", "Products, standards, execution, submittals, warranties", "Required manufacturer or concrete performance"],
            ["Addenda and RFIs", "Later clarifications or corrections", "Architect confirms which of two conflicting details governs"],
        ],
        [1950, 3450, 3960],
        compact=True,
    )

    add_heading(doc, "Trade and CSI division", 2)
    add_body(doc, "A trade is a package of work purchased from a specialist subcontractor. CSI numbers are standardized categories used to organize construction information. For example, the workbook uses 02 for demolition, 03 for concrete, 09 for finishes such as drywall and painting, 22 for plumbing, 23 for HVAC, and 26 for electrical.")
    add_callout(doc, "IMPORTANT", "The CSI number helps organize information, but it does not automatically decide contractual responsibility. Scope crossover still requires human review.")

    add_heading(doc, "Key terms before looking at Excel", 2)
    terms = [
        ("Scope of work", "The checklist or contract description of everything a trade must provide."),
        ("Proposal / bid", "The subcontractor's submitted price and commercial qualifications."),
        ("Takeoff", "A measured quantity from scaled drawings, such as 4,200 square feet of gym flooring."),
        ("RFI", "Request for Information: a formal question raised because information is missing, unclear, or conflicting."),
        ("Alternate", "An optional add or deduct price for a defined design choice."),
        ("Allowance", "A placeholder amount for work whose final details or quantity are not yet known."),
        ("Value engineering (VE)", "A proposal to reduce cost or improve value through an alternative material or method."),
        ("Exhibit A", "The trade-specific scope attached to the subcontract and made contractually binding."),
    ]
    add_table(doc, ["Term", "Meaning"], terms, [2400, 6960], compact=True)

    page_break(doc)

    add_heading(doc, "4. What every workbook sheet means", 1)
    add_body(doc, "The Jelleff workbook is best understood as a small database disguised as Excel. Project records, trade records, scope items, bidders, responses, budgets, and decisions are stored in cells and linked by formulas.")

    add_heading(doc, "The workbook has four kinds of sheets", 2)
    add_table(
        doc,
        ["Category", "Sheets", "Purpose"],
        [
            ["Project control", "Bid Sheet; MCN Summary", "Combines all trade costs into the project budget and summary."],
            ["Risk and coordination", "ScopeCrossOver; RFI Log; VEs; Alternates; SpecMatrix", "Tracks gaps, questions, options, and responsibility across trades."],
            ["Reusable structure", "SOWTemplate; Version History", "Provides the blank trade layout and tracks revisions."],
            ["Trade worksheets", "Demo, Conc, GWB, HVAC, Elect, Plumb, and dozens more", "Compares subcontractors for one trade at a time."],
        ],
        [2050, 3100, 4210],
    )

    add_heading(doc, "The important utility sheets", 2)
    utility_rows = [
        ["Bid Sheet", "Master list of trades, CSI code, assigned estimator, old/current budget, variance, cost per GSF, and bidder values.", "Output / control"],
        ["MCN Summary", "Rolls Bid Sheet values into standard CSI cost categories, total trades cost, construction cost, and cost per square foot.", "Calculated summary"],
        ["GCs n GRs", "General Conditions and General Requirements: project staff, supervision, temporary facilities, safety, administration, and similar GC costs.", "Inputs + formulas"],
        ["SOWTemplate", "Blank model for one trade: scope rows, internal budget, seven bidder groups, bond, totals, level state, and rank.", "Reusable template"],
        ["ScopeCrossOver", "Records work that could be missed or counted twice because several trades could reasonably own it.", "Human coordination"],
        ["RFI Log", "Stores question number, discipline, question, submission information, design response, and response date.", "Human workflow"],
        ["VEs", "Logs value-engineering ideas and cost impact.", "Human input"],
        ["Alternates", "Collects optional add/deduct prices by trade.", "Input + summary"],
        ["SpecMatrix", "Maps specification sections and bid packages to assigned estimators and coverage.", "Assignment / QA"],
    ]
    add_table(doc, ["Sheet", "Plain-English meaning", "Data type"], utility_rows, [1800, 6000, 1560], compact=True)

    add_heading(doc, "Most other tabs are trades", 2)
    add_body(doc, "Examples include Demo (demolition), Conc (concrete), Masonry, SS & MM (structural steel and miscellaneous metals), GWB (gypsum wallboard/drywall), DFH (doors, frames and hardware), Alum (aluminum/storefront), FireP (fire protection), Plumb, HVAC, Elect, LV (low voltage), EARTH, SiteConc, WetUtil, DryUtil, Lscape, and FFE.")

    page_break(doc)

    add_heading(doc, "5. A trade sheet explained cell by cell", 1)
    add_figure(doc, SHEET_IMAGE, "Figure 2. The anatomy of one trade sheet.", "Diagram showing project metadata at the top, master scope and budget fields on the left, repeating bidder groups on the right, and formula outputs at the bottom.")

    add_heading(doc, "The top of the sheet", 2)
    add_table(
        doc,
        ["Field", "Who supplies it", "What it means"],
        [
            ["Project, bid date, trade, revision", "Bid captain / project setup", "Identifies the bid event and package."],
            ["Estimator contact", "Project team", "Person responsible for this trade."],
            ["Bidder company and contact", "Proposal / bidder database", "The subcontractor being compared."],
            ["Level Scope?", "Estimator", "Whether review is complete enough to calculate/rank."],
            ["Base Bid", "Subcontractor proposal", "The bidder's submitted starting price."],
            ["Adjustments", "Estimator judgment", "Internal amount added or deducted to make scope comparable."],
            ["Bond", "Workbook formula", "In the Jelleff template, 2% of base bid plus adjustments."],
            ["Total Bid and Rank", "Workbook formula", "Comparable total and position among completed bidders."],
        ],
        [2100, 2550, 4710],
        compact=True,
    )

    add_heading(doc, "The left side: Marcel's requirement", 2)
    add_body(doc, "Each scope row describes one required item. The Desired Scope column records what the GC wants—usually Y. The response/comment and internal MCN Budget fields can hold quantity, unit of measure, unit price, and cost.")
    add_table(
        doc,
        ["Workbook field", "Meaning"],
        [
            ["Scope Items to include", "The master checklist the bidders must be compared against."],
            ["Desired Scope", "Whether the GC expects the item to be included."],
            ["Response / Comment", "Internal clarification or response associated with the requirement."],
            ["QTY", "Quantity."],
            ["UOM", "Unit of measure: SF, LF, CY, EA, LS, and similar units."],
            ["$/UOM", "Expected cost per unit."],
            ["COST", "Quantity multiplied by unit price."],
        ],
        [2500, 6860],
        compact=True,
    )

    add_heading(doc, "The right side: the bidders", 2)
    add_body(doc, "Every subcontractor receives the same three columns for every scope row. This is the heart of bid leveling.")
    add_table(
        doc,
        ["Field", "Meaning", "Example"],
        [
            ["Y", "Complete inclusion", "Proposal clearly includes removal and disposal."],
            ["N", "Excluded", "Proposal says concrete footings are excluded."],
            ["P", "Partial inclusion", "Includes one mobilization, but the scope requires several."],
            ["Note", "Reason and evidence", "“Make-safe by others” or proposal page reference."],
            ["Adjustment", "Internal cost needed to normalize the bid", "$18,000 added for a missing second mobilization."],
        ],
        [1100, 3700, 4560],
        compact=True,
    )
    add_callout(doc, "DO NOT CONFUSE THESE", "Desired Scope = what Marcel wants. Bidder Y/N/P = what a particular proposal appears to include. They answer different questions.")

    page_break(doc)
    add_heading(doc, "What the Jelleff demolition sheet contains", 2)
    actual_bids = [
        ["Celtic Demo", "$290,000", "FULL SCOPE", "$295,800"],
        ["Marc's Demo", "$240,000", "FULL SCOPE", "$244,800"],
        ["ACECO", "$358,597", "FULL SCOPE", "$365,769"],
        ["Berg Demo", "$154,826", "FULL SCOPE", "$157,923"],
        ["D&H", "$255,000", "FULL SCOPE", "$260,100"],
        ["Rath", "$260,000", "NO", "Not leveled"],
    ]
    add_table(doc, ["Bidder", "Base bid", "Workbook status", "Cached total"], actual_bids, [2650, 1900, 2100, 2710], compact=True)
    add_body(doc, "These values prove the workbook structure and formulas, but the file is not a complete truth set for automatic leveling: several bidder response columns are still blank, and the cached adjustment totals shown above are zero. Therefore, the lowest cached total should not be treated as an award recommendation without the completed scope review.")

    page_break(doc)

    add_heading(doc, "6. Bid leveling with a worked example", 1)
    add_body(doc, "Bid leveling means converting differently scoped proposals into comparable internal totals. The following example is deliberately simplified and is not taken from the Jelleff bidder results.")

    add_heading(doc, "The required demolition scope", 2)
    required = [
        ["1", "Demolish existing building structure", "Required"],
        ["2", "Remove foundations and footings", "Required"],
        ["3", "Haul and legally dispose of debris", "Required"],
    ]
    add_table(doc, ["Item", "Master scope", "Desired scope"], required, [950, 6510, 1900], compact=True)

    add_heading(doc, "Two proposals arrive", 2)
    comparison = [
        ["Base bid", "$290,000", "$240,000"],
        ["Building demolition", "Y", "Y"],
        ["Foundations and footings", "Y", "N — excluded"],
        ["Haul and disposal", "Y", "P — hauling only"],
        ["Internal adjustment", "$0", "+$50,000"],
        ["Subtotal before bond", "$290,000", "$290,000"],
        ["2% bond", "$5,800", "$5,800"],
        ["Leveled total", "$295,800", "$295,800"],
    ]
    table = add_table(doc, ["Comparison", "Bidder A", "Bidder B"], comparison, [3200, 3080, 3080], compact=True)
    for cell in table.rows[-1].cells:
        set_cell_shading(cell, PALE_BLUE)
        for p in cell.paragraphs:
            for run in p.runs:
                set_run_font(run, bold=True, color=NAVY)

    add_callout(doc, "THE LESSON", "Bidder B looked $50,000 cheaper, but only because important work was missing. After leveling, both bids cost the GC the same amount.")

    add_heading(doc, "The formula in Marcel's template", 2)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(6)
    p.paragraph_format.space_after = Pt(10)
    r = p.add_run("Leveled total = (Base bid + Approved adjustments) × 1.02")
    set_run_font(r, size=14, color=NAVY, bold=True)
    add_body(doc, "The 1.02 factor represents the workbook's 2% bond. Other companies or projects may use different rules, so the percentage must be a visible project setting—not hard-coded into the future product.")

    add_heading(doc, "Why the lowest bid does not automatically win", 2)
    for item in [
        "The low bidder may have exclusions or qualifications.",
        "The bidder may lack capacity, manpower, schedule availability, bonding, or relevant experience.",
        "The scope may still be ambiguous or dependent on an unanswered RFI.",
        "The estimator may prefer a slightly higher bidder with a cleaner and more reliable proposal.",
        "The final award must reflect scope-review and commercial negotiations, not only a rank formula.",
    ]:
        add_bullet(doc, item)

    page_break(doc)

    add_heading(doc, "7. What is manual, what Excel already automates, and what Structor should automate", 1)
    add_table(
        doc,
        ["Activity", "Today", "Future Structor behavior"],
        [
            ["Read drawings/specifications", "Manual specialist review; Wyre AI used as reference tool", "Draft trade scopes with page/detail/spec citations."],
            ["Write scope rows", "Estimator writes or edits Excel/PDF scope", "Suggest project-specific and library-based rows; human approves."],
            ["Enter bidder/contact/base bid", "Typed from email and proposal", "Extract from uploaded proposal or dedicated bid inbox."],
            ["Mark Y/N/P", "Estimator reads proposal row by row", "AI drafts status with exact proposal evidence; human confirms."],
            ["Write notes", "Manual", "Draft concise note from proposal exclusion or qualification."],
            ["Estimate adjustment", "Estimator judgment", "Suggest quantity × unit rate or historical allowance; require approval."],
            ["Bond, total, rank", "Excel formulas", "Deterministic calculation service using visible project rules."],
            ["Update Bid Sheet / summary", "Linked cells and manual review", "Automatic approved roll-up with audit trail."],
            ["Create Exhibit A", "Copy/rewrite approved scope", "Generate from approved scope and project template; legal review required."],
        ],
        [2400, 3000, 3960],
        compact=True,
    )

    add_heading(doc, "The key insight", 2)
    add_callout(doc, "THE PRODUCT IS NOT AN EXCEL CLONE", "Excel already performs the arithmetic. Structor's value is extracting, organizing, comparing, citing, reviewing, and reusing the information that humans currently type into Excel.")

    add_heading(doc, "What the current prototype does", 2)
    add_body(doc, "The prototype imports the Jelleff Demo or SOWTemplate tab, reads project/trade metadata, scope rows, bidder names, base bids, Y/N/P, notes and adjustments, calculates the 2% bond and leveled totals, supports review, and exports changes back into Excel. It demonstrates the workflow shell, not the finished AI intelligence.")
    if PROTOTYPE_IMAGE.exists():
        add_figure(doc, PROTOTYPE_IMAGE, "Figure 3. Current Structor workflow prototype using the Jelleff demolition structure. Some interface adjustment examples are illustrative, not source-of-truth workbook values.", "Desktop screenshot of the Structor prototype with project totals, demolition scope rows, bidder columns, Y/N/P statuses, notes, adjustments, and a selected-scope side panel.")

    add_heading(doc, "What the prototype does not yet do", 2)
    for item in [
        "Read construction drawings or specifications.",
        "Generate scope rows with citations.",
        "Parse real subcontractor proposal PDFs.",
        "Decide Y/N/P from evidence.",
        "Suggest defensible adjustment amounts from quantities and unit costs.",
        "Operate a dedicated bid inbox or subcontractor portal.",
        "Generate final contracts or takeoffs.",
    ]:
        add_bullet(doc, item)

    page_break(doc)

    add_heading(doc, "8. What you actually promised Marcel", 1)
    add_body(doc, "The meeting contains several ideas, but the immediate commitment was narrower than the long-term vision. It is important not to combine every requested feature into a one-week promise.")

    add_heading(doc, "The immediate commitment for the next meeting", 2)
    commitments = [
        ["Domain research", "Understand the terminology, document types, trade codes, and where information flows."],
        ["A basic front-end demo", "Show what the web application could look like so Marcel can give workflow feedback."],
        ["A sheet-filling application", "Prototype an application that populates the existing trade-sheet structure."],
        ["Excel export", "Preserve Marcel's familiar workbook output."],
        ["Iterative review", "Meet weekly, ask informed questions, and refine the interface and workflow."],
    ]
    add_table(doc, ["Commitment", "Meaning"], commitments, [2700, 6660], compact=True)

    add_callout(doc, "EXACT BOUNDARY", "You estimated that the sheet prototype and demo could be ready by the following meeting. You did not promise a production-ready drawing-reading AI, automatic takeoffs, or a fully autonomous estimator within one week.")

    add_heading(doc, "The broader capability you said was possible", 2)
    add_body(doc, "You also agreed that the product could eventually extract construction information from blueprints, organize it by universal construction codes, show scopes as notes with source references, and export a trade file. That was agreement on technical direction, not a one-week completion date.")

    add_heading(doc, "There were two different 'Step 1' statements", 2)
    add_table(
        doc,
        ["Whose Step 1?", "What it meant", "Correct interpretation"],
        [
            ["Marcel's product Step 1", "Upload drawings/specifications and generate cited scopes by trade.", "The first major AI capability and the hardest intelligence problem."],
            ["The team's implementation starting point", "Build the interface and automate sheet population/export.", "The easiest demonstrable prototype and the actual next-meeting commitment."],
        ],
        [2200, 3700, 3460],
    )

    add_heading(doc, "What 'automate the sheets' should mean", 2)
    for item in [
        "Create or populate a trade sheet from structured project, scope, and bidder data.",
        "Accept scope rows from manual entry first, then from the drawing reader later.",
        "Accept bidder data from manual entry first, then proposal/email extraction later.",
        "Calculate adjustments, bond, leveled total, completion state, and rank consistently.",
        "Update the master Bid Sheet only after human approval.",
        "Export a workbook that preserves the expected layout and opens without repair warnings.",
    ]:
        add_bullet(doc, item)

    page_break(doc)

    add_heading(doc, "9. The recommended product milestones and data model", 1)
    add_heading(doc, "Milestone 1 — Workflow and Excel prototype", 2)
    add_body(doc, "This is the current stage: project setup, trade selection, structured scope rows, bidder values, review controls, deterministic leveling calculations, and Excel export. It proves that the product can reproduce Marcel's operating surface without waiting for complex document AI.")

    add_heading(doc, "Milestone 2 — First true AI scope pilot", 2)
    add_body(doc, "Use one bounded trade—recommended: Jelleff drywall. Upload the relevant drawing/specification subset, generate a draft scope, attach a citation to every project-specific item, let Marcel accept/edit/reject it, and export the approved result.")

    add_heading(doc, "Milestone 3 — Proposal ingestion and auto-leveling", 2)
    add_body(doc, "Parse a real bidder proposal, identify the project/trade/bidder, extract base price, exclusions, alternates and qualifications, compare the proposal with the approved scope, and draft Y/N/P, notes, and missing-scope flags for approval.")

    add_heading(doc, "Milestone 4 — Roll-up, contract, and operations", 2)
    add_body(doc, "Move approved values into the project budget, generate Exhibit A from the approved scope, and add RFI, VE, alternate, assignment, audit, and handoff workflows. Takeoffs and subcontractor-side pricing can follow as separate, measurable products.")

    add_heading(doc, "The data model behind Excel", 2)
    data_rows = [
        ["Project", "Name, client, location, dates, area, revision, source documents"],
        ["Trade", "CSI code, trade name, assigned estimator, project budget"],
        ["Scope item", "Section, text, required flag, source citation, review status"],
        ["Bidder", "Company, contact, email, phone, qualification data"],
        ["Proposal", "File, bidder, trade, submitted price, alternates, exclusions"],
        ["Bid response", "Scope item × bidder: Y/N/P, note, evidence, adjustment, approver"],
        ["Leveling result", "Base bid, adjustment total, bond, leveled total, completeness, rank"],
        ["Decision", "Recommended/selected bidder, approved budget value, rationale"],
        ["Award document", "Approved scope converted to Exhibit A / contract language"],
    ]
    add_table(doc, ["Record", "Important fields"], data_rows, [2200, 7160], compact=True)

    add_callout(doc, "WHY THIS MATTERS", "Once these records exist separately from Excel, the same approved data can power the web app, Excel export, project summary, scope library, audit history, and contract without copying information repeatedly.")

    page_break(doc)

    add_heading(doc, "10. How the AI components should work safely", 1)
    add_heading(doc, "A. Drawing and specification reader", 2)
    for number, item in enumerate([
        "Ingest each file and preserve page, sheet number, discipline, revision, and document type.",
        "Extract text, schedules, notes, callouts, details, and specification sections while retaining coordinates.",
        "Classify facts into likely trades and CSI sections.",
        "Generate draft scope sentences from the evidence and reusable scope library.",
        "Attach the exact drawing/specification citation and preview region to every project-specific scope item.",
        "Require a reviewer to accept, edit, reject, or reassign each item.",
    ], start=1):
        add_manual_number(doc, item, number)

    add_heading(doc, "B. Proposal reader", 2)
    for number, item in enumerate([
        "Identify project, trade, bidder, proposal date and revision.",
        "Extract base bid, taxes/bond assumptions, exclusions, alternates, allowances and qualifications.",
        "Compare proposal language to each approved master scope row.",
        "Draft Y/N/P and a concise note with the exact proposal evidence.",
        "Flag unresolved or contradictory language instead of guessing.",
        "Let the estimator approve the comparison before it affects totals or the budget.",
    ], start=1):
        add_manual_number(doc, item, number)

    add_heading(doc, "C. Adjustment suggestion", 2)
    add_body(doc, "The system can suggest an adjustment when a bidder misses work, but this must be explainable. A useful suggestion shows the missing quantity, unit of measure, selected unit rate, rate source, calculation, confidence, and who approved it. A black-box number is not acceptable for a bid or contract.")

    add_heading(doc, "D. Scope and cost library", 2)
    add_body(doc, "The library should remember approved, reusable knowledge—not silently train on every draft. Store standard scope clauses by trade, project-specific approved clauses, drawing/specification citation patterns, historical unit costs with date/location context, proposal terminology, and approved adjustment decisions. Every reuse should show provenance.")

    add_heading(doc, "Non-negotiable safeguards", 2)
    safeguards = [
        ["Evidence", "Every non-generic extracted fact links back to its source."],
        ["Human approval", "AI suggestions do not automatically alter the budget or contract."],
        ["Revision awareness", "A drawing/spec revision can invalidate earlier citations and decisions."],
        ["Audit trail", "Store who changed Y/N/P, adjustment, selection, and contract scope, and when."],
        ["Clean-room implementation", "Do not copy Wyre AI code, private prompts, protected templates, or NDA-restricted behavior."],
        ["Security", "Drawings, bids and contracts are commercially sensitive and require tenant isolation and access controls."],
    ]
    add_table(doc, ["Safeguard", "Required behavior"], safeguards, [2200, 7160], compact=True)

    page_break(doc)

    add_heading(doc, "11. What you need from Marcel next", 1)
    add_callout(doc, "DO NOT ASK FOR RANDOM MOCK SHEETS", "Ask for one matched, completed trade package from source documents through final decision. The files must tell one coherent story.")

    add_heading(doc, "The recommended golden dataset", 2)
    golden = [
        "The exact drawing and specification files used for the bid.",
        "The final human-written trade scope or bid form.",
        "At least two or three subcontractor proposals for that trade.",
        "The completed leveling sheet with Y/N/P, notes and adjustments.",
        "RFIs, addenda or clarifications that changed the scope.",
        "The selected bidder and approved budget value.",
        "The final Exhibit A, if one was issued.",
    ]
    for item in golden:
        add_bullet(doc, item)

    add_heading(doc, "Best first pilot: Jelleff drywall", 2)
    add_body(doc, "The verified source inventory indicates that drywall is the strongest complete test case: drawings and project manual, a human drywall scope/bid form, a District bid form and proposal, the populated GWB sheet, and a Wyre AI sample package. Demolition is useful for understanding workbook formulas, but the shared drywall package is better for testing the whole chain.")

    add_heading(doc, "What you should ask Marcel to label", 2)
    labels = [
        ["For each scope item", "Which drawing page/detail or specification section created it? Is it standard or project-specific?"],
        ["For trade ownership", "Why does this item belong to drywall/concrete/etc. and not another trade?"],
        ["For each proposal", "Which sentence supports Y, N or P? What does ambiguous wording mean?"],
        ["For each adjustment", "What quantity, unit rate and judgment created the amount?"],
        ["For the decision", "Why was the preferred bidder selected? What non-price factors mattered?"],
        ["For the output", "Which values moved to the Bid Sheet, budget and Exhibit A?"],
    ]
    add_table(doc, ["Labeling task", "Question to answer"], labels, [2300, 7060], compact=True)

    add_heading(doc, "Your role in the automation", 2)
    add_body(doc, "You do not need to become a senior estimator before building anything. Your job is to make the decision process explicit: identify the input, capture the output, record the evidence and approval, and translate repetitive actions into software. Marcel remains the domain owner who validates whether the output is commercially correct.")

    page_break(doc)

    add_heading(doc, "12. Questions for the next meeting", 1)
    questions = [
        "Which exact project and trade are we using as the first acceptance test?",
        "Is the next deliverable sheet population only, or must it parse a proposal PDF?",
        "Which workbook is the production template, and which formulas/formatting must be preserved exactly?",
        "What does FULL SCOPE mean operationally, and who is allowed to set it?",
        "Can Y/N/P be entered by the subcontractor, the estimator, AI, or all three?",
        "What evidence is sufficient to mark Y? What conditions require P instead?",
        "How are missing-scope adjustments calculated today?",
        "Does bond always apply to base bid plus adjustments, and is 2% configurable?",
        "When are allowances and alternates included in the leveled total?",
        "Which bidder value is copied into the Bid Sheet: lowest leveled, selected bidder, or manual budget?",
        "What must be completed before a bidder can be ranked?",
        "Which proposal formats and email patterns are most common?",
        "What may legally be used from Wyre AI and the prior development team?",
        "Who approves scopes, adjustments, budget values and Exhibit A?",
        "What measurable accuracy would make the first AI pilot successful?",
    ]
    for number, q in enumerate(questions, start=1):
        add_manual_number(doc, q, number)

    add_heading(doc, "A 30-second explanation you can repeat", 2)
    add_callout(doc, "SCRIPT", "Marcel receives construction drawings and needs to buy dozens of specialist trades. He writes a checklist for each trade, receives several proposals, and compares every proposal against the same checklist so a cheap bid cannot hide missing work. Today that comparison is entered manually into Excel. We are building a web app that structures the data, helps read drawings and proposals, drafts the comparison with evidence, calculates the leveled totals, and exports the approved result back to his workbook and eventually the contract.")

    add_heading(doc, "Glossary", 1)
    glossary = [
        ["Addendum", "A formal change or clarification issued during bidding."],
        ["Allowance", "A placeholder amount for uncertain or not-yet-defined work."],
        ["Alternate", "A separately priced optional add or deduct."],
        ["Base bid", "The subcontractor's primary submitted price before internal leveling adjustments."],
        ["Bid captain", "The person coordinating the overall project estimate and bid process."],
        ["Bid leveling", "Making proposals comparable against the same required scope."],
        ["Bond", "Financial security cost applied according to project/company rules; 2% in the Jelleff template."],
        ["CSI division", "Standard construction category used to organize work and specifications."],
        ["Exclusion", "Work the bidder explicitly says is not included."],
        ["Exhibit A", "The subcontract's trade-specific scope attachment."],
        ["GC", "General contractor."],
        ["GCs / GRs", "General Conditions and General Requirements: the GC's project-management and jobsite costs."],
        ["GSF", "Gross square feet of the project."],
        ["Leveled total", "Comparable internal total after approved adjustments and bond rules."],
        ["Proposal qualification", "A condition or assumption limiting the bidder's offer."],
        ["RFI", "Request for Information."],
        ["Scope crossover", "Work whose trade ownership is ambiguous or shared."],
        ["Scope item", "One requirement in a trade's master checklist."],
        ["Subcontractor", "Specialist company performing one part of the construction."],
        ["Takeoff", "Measured quantity derived from scaled drawings."],
        ["Trade", "A package of related construction work purchased from a specialist."],
        ["UOM", "Unit of measure, such as SF, LF, CY, EA or LS."],
        ["VE", "Value engineering."],
        ["Y/N/P", "Included / not included / partially included."],
    ]
    add_table(doc, ["Term", "Plain-English definition"], glossary, [2200, 7160], compact=True)

    add_heading(doc, "Sources and limitations", 1)
    add_body(doc, "This guide was prepared from the timestamped Zoom transcript for the 26 August 2026 discovery meeting, the Jelleff Rec Live Bid sheet.xlsx workbook, the verified shared-source inventory, and the current Structor prototype documentation. Workbook formulas and cached values were inspected directly.")
    add_body(doc, "The local MP4 recording could not be opened through the available browser because local-file playback was blocked by security policy, and the in-app Drive session required sign-in. Therefore, exact claims about the meeting are based on the complete timestamped transcript rather than unverified visual inference. The workbook and locally retained project materials were inspected directly.")
    add_body(doc, "Some transcript speaker labels appear imperfect, especially where a statement addressed to Marcel is labeled as Marcel speaking. Requirements are attributed by substance rather than assuming every speaker tag is exact.")

    source_rows = [
        ["Zoom transcript", "Meeting statements, terminology, requested product capabilities, and commitments; 18:02–18:38."],
        ["Jelleff Rec Live Bid sheet.xlsx", "Workbook sheets, formulas, bidder structure, budget roll-ups, RFI log, and trade tabs."],
        ["SOURCE_MATERIALS.md", "Verified shared-folder inventory dated 31 August 2026 and recommended golden dataset."],
        ["README.md", "Current prototype boundary, implementation status, roadmap, and source caveats."],
    ]
    add_table(doc, ["Source", "Used for"], source_rows, [2800, 6560], compact=True)

    # Core properties and final save
    doc.core_properties.title = "Understanding Structor AI"
    doc.core_properties.subject = "Beginner's guide to Marcel Khoury's preconstruction workflow and the promised automation"
    doc.core_properties.author = "Synix Solutions"
    doc.core_properties.keywords = "preconstruction, bid leveling, scope of work, Structor AI, Jelleff"
    doc.save(DOCX_PATH)
    print(DOCX_PATH)


if __name__ == "__main__":
    build_document()
